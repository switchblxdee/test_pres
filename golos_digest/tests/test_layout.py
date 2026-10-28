import random
import re
import zipfile

import pytest
from lxml import etree
from pptx import Presentation

from golos_digest.data import load_digest
from golos_digest.agent import build
from golos_digest.measure import Measurer
from golos_digest.quotes import QuoteSelector
from golos_digest.render import (
    BOTTOM, CARD_PAD, ROW_OFFSET, TOP_FIRST, TOP_OTHER, Layouter,
)
from tests.conftest import markup_row, summary_row

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
RPR_ORDER = ["ln", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill", "effectLst", "effectDag",
             "highlight", "uLnTx", "uLn", "uFillTx", "uFill", "latin", "ea", "cs", "sym", "hlinkClick",
             "hlinkMouseOver", "rtl", "extLst"]
PPR_ORDER = ["lnSpc", "spcBef", "spcAft", "buClrTx", "buClr", "buSzTx", "buSzPct", "buSzPts", "buFontTx", "buFont",
             "buNone", "buAutoNum", "buChar", "buBlip", "tabLst", "defRPr", "extLst"]

WORDS = ("сервис падает ошибка таймаут доступ модель генерация MCP токен авторизация интеграция "
         "релиз версия CLI Desktop пайплайн https://example.local/very/long/path/to/some/page?id=123456 "
         "нагрузка индексация репозиторий").split()


def rand_text(rng, lo, hi):
    return " ".join(rng.choice(WORDS) for _ in range(rng.randint(lo, hi))).capitalize() + "."


def dataset(seed, n_products=4, max_topics=9):
    rng = random.Random(seed)
    s, m = [], []
    for p in range(n_products):
        product = f"Продукт {p}" if p else "Очень длинное название продукта " * 4
        for k in range(rng.randint(1, max_topics)):
            topic = rand_text(rng, 3, 40)
            prev = rng.choice([0, 0, 5, 10, 40, None])
            status = rng.choice(["", "в списке болей", "кандидат в боль - на анализе", rand_text(rng, 2, 60),
                                 "возобновилось"])
            s.append(summary_row(product, topic, rng.randint(1, 200), prev, status=status))
            for _ in range(rng.randint(0, 6)):
                m.append(markup_row(product, topic, rand_text(rng, 2, 120), day=rng.randint(9, 15)))
    return s, m


def check_plan(plans):
    for i, sl in enumerate(plans):
        assert sl.sections, f"пустой слайд {i}"
        start = TOP_FIRST if sl.first else TOP_OTHER
        prev_bottom = start
        for sec in sl.sections:
            assert sec.top >= prev_bottom, "карточки продуктов наложились"
            assert sec.bottom <= BOTTOM, "карточка вылезла за низ слайда"
            row_floor = sec.top + ROW_OFFSET
            for r in sec.rows:
                assert r.top >= row_floor, "строки наложились"
                assert r.top + r.height + CARD_PAD <= BOTTOM + 1
                assert r.box_h >= r.text_h, "карточка статуса ниже текста строки"
                row_floor = r.top + r.height
            prev_bottom = sec.bottom


def check_package(path):
    z = zipfile.ZipFile(path)
    names = z.namelist()
    assert not any("thumbnail" in n for n in names)
    assert sum("slideLayouts/slideLayout" in n and n.endswith(".xml") for n in names) == 1
    for n in names:
        if not re.match(r"ppt/slides/slide\d+\.xml$", n):
            continue
        t = etree.fromstring(z.read(n))
        ids = t.xpath("//p:cNvPr/@id", namespaces={"p": "http://schemas.openxmlformats.org/presentationml/2006/main"})
        assert len(ids) == len(set(ids)), f"дубли id в {n}"
        for tag, order in (("rPr", RPR_ORDER), ("endParaRPr", RPR_ORDER), ("pPr", PPR_ORDER)):
            for el in t.iter(A + tag):
                idx = [order.index(c.tag.replace(A, "")) for c in el if c.tag.replace(A, "") in order]
                assert idx == sorted(idx), f"порядок детей {tag} в {n}"
    prs = Presentation(path)
    for slide in prs.slides:
        for blip in slide.shapes._spTree.xpath(".//a:blip/@r:embed"):
            assert blip in slide.part.rels
    assert prs.core_properties.author == "Голос IT"


@pytest.mark.parametrize("seed", range(6))
def test_layout_invariants_random(make_xlsx, cfg, tmp_path, seed):
    s, m = dataset(seed)
    path = make_xlsx(s, m, name=f"r{seed}.xlsx")
    d = load_digest(path, cfg)
    selector = QuoteSelector(cfg)
    for t in d.topics:
        t.quotes = selector.select(t) if t.candidates else []
    plans = Layouter(cfg, Measurer(cfg.fonts)).plan(d)
    check_plan(plans)
    assert sum(len(sec.rows) for sl in plans for sec in sl.sections) == len(d.topics)


def test_end_to_end_file(make_xlsx, cfg, tmp_path):
    s, m = dataset(42, n_products=3)
    out = build(make_xlsx(s, m), tmp_path / "out.pptx", cfg)
    check_package(out)
    assert out.stat().st_size < 300_000  # без мёртвых фонов шаблона


def test_long_texts_are_clipped(make_xlsx, cfg):
    long = "слово " * 400
    s = [summary_row("P" * 300, long, 10, 5, status=long)]
    m = [markup_row("P" * 300, long, long)]
    d = load_digest(make_xlsx(s, m), cfg)
    d.topics[0].quotes = [long]
    lay = Layouter(cfg, Measurer(cfg.fonts))
    row = lay.measure_row(1, d.topics[0])
    assert row.title.endswith("…") and row.status.endswith("…") and row.quotes[0].endswith("…»")
    assert row.height <= lay.max_row_h


@pytest.mark.parametrize("seed", range(10, 16))
def test_global_plan_never_worse_than_greedy(make_xlsx, cfg, seed):
    s, m = dataset(seed)
    d = load_digest(make_xlsx(s, m, name=f"g{seed}.xlsx"), cfg)
    selector = QuoteSelector(cfg)
    for t in d.topics:
        t.quotes = selector.select(t) if t.candidates else []
    lay = Layouter(cfg, Measurer(cfg.fonts))
    best = lay.plan(d)
    check_plan(best)
    assert len(best) <= len(lay._plan(d, cfg.quotes.max_quote_lines))


def test_kpi_labels_fit_wide_fallback_fonts(make_xlsx, cfg, tmp_path):
    """«проанализировано» должно влезать в строку даже шрифтом DejaVu Sans (самый широкий
    из типовых заменителей SB Sans): место под текст ≥ 1.40", подпись не вылезает за карточку."""
    from golos_digest.ooxml import get_box, xp
    out = build(make_xlsx([summary_row("P", "T", 5, 4)], [markup_row("P", "T", "сообщение")]),
                tmp_path / "k.pptx", cfg)
    tree = Presentation(out).slides[0].shapes._spTree
    for group in tree.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}grpSp"):
        names = {c.get("name") for c in group.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")}
        if not any(n and n.startswith("TextBox") for n in names):
            continue
        for sp in group.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}sp"):
            txt = "".join(xp(sp, ".//a:t/text()"))
            if "проанализировано" in txt:
                card = [s for s in group.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}sp")
                        if xp(s, "./p:spPr/a:prstGeom[@prst='roundRect']") and get_box(s)[2] > 1_500_000][0]
                lx, _, lw, _ = get_box(sp)
                cx, _, cw, _ = get_box(card)
                assert xp(sp, ".//a:bodyPr/@lIns") == ["0"]
                assert lw >= int(1.40 * 914400)
                assert lx + lw <= cx + cw
                return
    pytest.fail("подпись KPI не найдена")
